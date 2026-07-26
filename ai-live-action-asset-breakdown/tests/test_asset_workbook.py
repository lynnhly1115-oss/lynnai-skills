from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

SKILL_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL_ROOT / "scripts"))

import build_asset_workbook as builder


class AssetWorkbookTests(unittest.TestCase):
    def sample_data(self) -> dict:
        return {
            "projectTitle": "《测试项目》",
            "episodes": [
                {
                    "episode": "第1集",
                    "assets": [
                        {
                            "scene": "1-1",
                            "category": "人物",
                            "name": "沈晚晴",
                            "assetType": "基础",
                            "parent": "",
                            "description": "主要角色，剧本未明确服装",
                        },
                        {
                            "scene": "1-2",
                            "category": "人物",
                            "name": "沈晚晴·治安局制服",
                            "assetType": "变体",
                            "parent": "沈晚晴",
                            "description": "更换治安局制服",
                        },
                        {
                            "scene": "1-1",
                            "category": "场景",
                            "name": "百乐门·舞厅",
                            "assetType": "基础",
                            "parent": "",
                            "description": "舞厅内景",
                        },
                        {
                            "scene": "1-1",
                            "category": "道具",
                            "name": "口红",
                            "assetType": "基础",
                            "parent": "",
                            "description": "沈晚晴使用的口红",
                        },
                    ],
                },
                {
                    "episode": "第2集",
                    "assets": [
                        {
                            "scene": "2-3",
                            "category": "场景",
                            "name": "贫民窟·暗巷",
                            "assetType": "基础",
                            "parent": "",
                            "description": "夜间暗巷",
                        }
                    ],
                },
            ],
        }

    def test_builds_single_sheet_with_locked_contract(self) -> None:
        data = builder.validate_and_normalize(self.sample_data())
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "assets.xlsx"
            result = builder.build_workbook(data, output)
            inspection = builder.inspect_workbook(output, result["assetRows"])
            workbook = load_workbook(output, data_only=False)
            sheet = workbook["资产统计表"]
            headers = [sheet.cell(4, col).value for col in range(1, 14)]
            gray_fill = sheet["K5"].fill.fgColor.rgb
            scene_fill = sheet["K7"].fill.fgColor.rgb
            formula = sheet["G9"].value
            workbook.close()

        self.assertTrue(inspection["passed"], inspection["issues"])
        self.assertEqual(headers, builder.HEADERS)
        self.assertEqual(inspection["assetRows"], 5)
        self.assertEqual(inspection["freezePanes"], "H5")
        self.assertEqual(gray_fill, "00ECEBE8")
        self.assertEqual(scene_fill, "00FFFFFF")
        self.assertTrue(str(formula).startswith("="))

    def test_rejects_variant_before_its_parent(self) -> None:
        data = self.sample_data()
        data["episodes"][0]["assets"] = [data["episodes"][0]["assets"][1]]
        with self.assertRaisesRegex(ValueError, "missing or later parent"):
            builder.validate_and_normalize(data)

    def test_extracts_docx_paragraphs_and_tables(self) -> None:
        document = Document()
        document.add_paragraph("第一集")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "1-1 百乐门舞厅 夜 内"
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "script.docx"
            document.save(source)
            text = builder.extract_script_text(source)
        self.assertIn("第一集", text)
        self.assertIn("百乐门舞厅", text)

    def test_writes_local_html_preview(self) -> None:
        data = builder.validate_and_normalize(self.sample_data())
        with tempfile.TemporaryDirectory() as temp_dir:
            preview = builder.write_html_preview(data, Path(temp_dir))
            content = preview.read_text(encoding="utf-8")
        self.assertIn("资产统计表预览", content)
        self.assertIn("沈晚晴·治安局制服", content)


if __name__ == "__main__":
    unittest.main()
