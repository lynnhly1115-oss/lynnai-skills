#!/usr/bin/env python3
from __future__ import annotations

import argparse
import html
import json
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from dependency_bootstrap import ensure_dependencies

ensure_dependencies(
    "ai-live-action-asset-breakdown",
    Path(__file__).with_name("requirements.txt"),
    ("docx", "openpyxl", "pdfplumber"),
)

import pdfplumber
from docx import Document
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

HEADERS = [
    "集数",
    "场次",
    "类别",
    "资产名称",
    "基础/变体",
    "母资产",
    "资产说明",
    "图片1",
    "图片2",
    "图片3",
    "图片4",
    "图片5",
    "图片6",
]
CATEGORY_ORDER = {"人物": 0, "场景": 1, "道具": 2}
CATEGORY_FILLS = {"人物": "F7E8E3", "场景": "F5E9CB", "道具": "E9F0E8"}


def natural_key(value: str) -> list[Any]:
    return [int(part) if part.isdigit() else part.casefold() for part in re.split(r"(\d+)", value)]


def validate_and_normalize(data: Any) -> dict[str, Any]:
    if not isinstance(data, dict):
        raise TypeError("Input JSON must be an object")
    project_title = data.get("projectTitle")
    if not isinstance(project_title, str) or not project_title.strip():
        raise ValueError("projectTitle is required")
    episode_blocks = data.get("episodes")
    if not isinstance(episode_blocks, list) or not episode_blocks:
        raise ValueError("episodes must be a non-empty array")

    known_names: set[str] = set()
    episodes: list[dict[str, Any]] = []
    for episode_index, block in enumerate(episode_blocks, start=1):
        if not isinstance(block, dict):
            raise TypeError(f"Episode {episode_index} must be an object")
        episode = block.get("episode")
        assets = block.get("assets")
        if not isinstance(episode, str) or not episode.strip() or not isinstance(assets, list):
            raise ValueError(f"Episode {episode_index} must contain episode and assets")

        normalized_assets: list[dict[str, str | int]] = []
        for asset_index, asset in enumerate(assets, start=1):
            if not isinstance(asset, dict):
                raise TypeError(f"{episode} asset {asset_index} must be an object")
            clean: dict[str, str | int] = {"sourceOrder": asset_index}
            for key in ("scene", "category", "name", "assetType", "description"):
                value = asset.get(key)
                if not isinstance(value, str) or not value.strip():
                    raise ValueError(f"{episode} asset {asset_index}: {key} is required")
                clean[key] = value.strip()
            category = str(clean["category"])
            asset_type = str(clean["assetType"])
            if category not in CATEGORY_ORDER:
                raise ValueError(f"Invalid category: {category}")
            if asset_type not in {"基础", "变体"}:
                raise ValueError(f"Invalid assetType: {asset_type}")
            parent = asset.get("parent", "")
            parent = parent.strip() if isinstance(parent, str) else ""
            name = str(clean["name"])
            if asset_type == "变体" and not parent:
                raise ValueError(f"Variant {name} requires parent")
            if asset_type == "基础" and parent:
                raise ValueError(f"Base asset {name} must not have parent")
            if name in known_names:
                raise ValueError(f"Duplicate asset name: {name}")
            if asset_type == "变体" and parent not in known_names:
                raise ValueError(f"Variant {name} references missing or later parent: {parent}")
            known_names.add(name)
            clean["parent"] = parent
            normalized_assets.append(clean)

        normalized_assets.sort(
            key=lambda item: (
                CATEGORY_ORDER[str(item["category"])],
                natural_key(str(item["scene"])),
                int(item["sourceOrder"]),
            )
        )
        episodes.append({"episode": episode.strip(), "assets": normalized_assets})
    return {"projectTitle": project_title.strip(), "episodes": episodes}


def iter_docx_lines(path: Path) -> list[str]:
    document = Document(path)
    lines: list[str] = []
    if hasattr(document, "iter_inner_content"):
        blocks = document.iter_inner_content()
    else:  # pragma: no cover - compatibility with older python-docx
        blocks = []
    for block in blocks:
        if isinstance(block, Paragraph):
            if block.text.strip():
                lines.append(block.text.strip())
        elif isinstance(block, DocxTable):
            for row in block.rows:
                for cell in row.cells:
                    text = "\n".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
                    if text:
                        lines.append(text)
    if not lines:
        lines.extend(paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip())
    return lines


def convert_legacy_doc(path: Path, output_dir: Path) -> Path:
    converted = output_dir / f"{path.stem}.docx"
    if shutil.which("textutil"):
        command = ["textutil", "-convert", "docx", "-output", str(converted), str(path)]
    else:
        office = shutil.which("soffice") or shutil.which("libreoffice")
        if office is None:
            raise RuntimeError(
                "Legacy .doc conversion requires macOS textutil or LibreOffice. "
                "Install LibreOffice or export the script as .docx."
            )
        command = [office, "--headless", "--convert-to", "docx", "--outdir", str(output_dir), str(path)]
    subprocess.run(command, check=True, capture_output=True, text=True)
    if converted.exists():
        return converted
    candidates = list(output_dir.glob("*.docx"))
    if len(candidates) != 1:
        raise RuntimeError("Legacy .doc conversion did not produce a readable .docx file")
    return candidates[0]


def extract_script_text(source: Path) -> str:
    suffix = source.suffix.lower()
    if suffix in {".txt", ".md"}:
        return source.read_text(encoding="utf-8-sig")
    if suffix in {".docx", ".docm"}:
        return "\n".join(iter_docx_lines(source))
    if suffix == ".doc":
        with tempfile.TemporaryDirectory(prefix="asset-doc-") as temp_dir:
            return "\n".join(iter_docx_lines(convert_legacy_doc(source, Path(temp_dir))))
    if suffix == ".pdf":
        pages: list[str] = []
        with pdfplumber.open(source) as pdf:
            for page in pdf.pages:
                text = page.extract_text(x_tolerance=2, y_tolerance=3) or ""
                if text.strip():
                    pages.append(text.strip())
        if not pages:
            raise RuntimeError("The PDF contains no extractable text; use OCR or provide a Word/text export")
        return "\n\n".join(pages)
    raise ValueError(f"Unsupported input format: {source.suffix}")


def make_border(color: str = "D9D1C7", style: str = "thin") -> Border:
    side = Side(style=style, color=color)
    return Border(left=side, right=side, top=side, bottom=side)


def asset_rows(data: dict[str, Any]) -> tuple[list[list[Any]], list[dict[str, str]]]:
    rows: list[list[Any]] = []
    meta: list[dict[str, str]] = []
    for block in data["episodes"]:
        for asset in block["assets"]:
            rows.append(
                [
                    block["episode"],
                    asset["scene"],
                    asset["category"],
                    asset["name"],
                    asset["assetType"],
                    asset["parent"],
                    asset["description"],
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                ]
            )
            meta.append({"kind": "asset", "category": asset["category"], "assetType": asset["assetType"]})
        rows.append([block["episode"], "", "小计", "本集小计", "", "", None, None, None, None, None, None, None])
        meta.append({"kind": "subtotal"})
    rows.append(["全剧", "", "总计", "累计唯一资产", "", "", None, None, None, None, None, None, None])
    meta.append({"kind": "grandTotal"})
    return rows, meta


def build_workbook(data: dict[str, Any], output: Path) -> dict[str, Any]:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "资产统计表"
    sheet.sheet_view.showGridLines = False

    sheet.merge_cells("A1:M1")
    sheet["A1"] = f"{data['projectTitle']} AI真人剧资产统计表"
    sheet.merge_cells("A2:M2")
    sheet["A2"] = "统计规则：每项资产只在首次出现的集数中列出；后续复用不重复。服装、年龄、伪装、受伤、损毁与功能/特效变化按变体统计。"
    sheet.merge_cells("A3:M3")
    sheet["A3"] = "图片位：人物和戏用道具使用图片1-3；场景使用图片1-6。场景陈设并入场景说明，纯背景群众不单独统计。"
    sheet.append(HEADERS)

    rows, meta = asset_rows(data)
    for row in rows:
        sheet.append(row)
    start_row = 5
    end_row = 4 + len(rows)

    for offset, item in enumerate(meta):
        row = start_row + offset
        if item["kind"] == "subtotal":
            sheet.cell(row, 7).value = (
                f'=\"新增人物 \"&COUNTIFS($A${start_row}:$A${end_row},A{row},$C${start_row}:$C${end_row},\"人物\")'
                f'&\"｜新增场景 \"&COUNTIFS($A${start_row}:$A${end_row},A{row},$C${start_row}:$C${end_row},\"场景\")'
                f'&\"｜新增道具 \"&COUNTIFS($A${start_row}:$A${end_row},A{row},$C${start_row}:$C${end_row},\"道具\")'
                f'&\"｜合计 \"&(COUNTIFS($A${start_row}:$A${end_row},A{row},$C${start_row}:$C${end_row},\"人物\")'
                f'+COUNTIFS($A${start_row}:$A${end_row},A{row},$C${start_row}:$C${end_row},\"场景\")'
                f'+COUNTIFS($A${start_row}:$A${end_row},A{row},$C${start_row}:$C${end_row},\"道具\"))'
            )
        elif item["kind"] == "grandTotal":
            sheet.cell(row, 7).value = (
                f'=\"累计人物 \"&COUNTIF($C${start_row}:$C${end_row},\"人物\")'
                f'&\"｜累计场景 \"&COUNTIF($C${start_row}:$C${end_row},\"场景\")'
                f'&\"｜累计道具 \"&COUNTIF($C${start_row}:$C${end_row},\"道具\")'
                f'&\"｜总计 \"&(COUNTIF($C${start_row}:$C${end_row},\"人物\")'
                f'+COUNTIF($C${start_row}:$C${end_row},\"场景\")'
                f'+COUNTIF($C${start_row}:$C${end_row},\"道具\"))'
            )

    gold = "F1D39B"
    burgundy = "351719"
    sheet["A1"].font = Font(name="Microsoft YaHei", bold=True, color=gold, size=18)
    sheet["A1"].fill = PatternFill("solid", fgColor=burgundy)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 34
    for row, fill, color, height in ((2, "F3E4C2", "5A2426", 34), (3, "FBF7EF", "4C4A46", 30)):
        cell = sheet.cell(row, 1)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(name="Microsoft YaHei", color=color, size=10)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        sheet.row_dimensions[row].height = height
    for cell in sheet[4]:
        cell.fill = PatternFill("solid", fgColor="7A2E2E")
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF", size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = make_border("D9C7B8")
    sheet.row_dimensions[4].height = 34

    for offset, item in enumerate(meta):
        row = start_row + offset
        if item["kind"] == "asset":
            fill = PatternFill("solid", fgColor=CATEGORY_FILLS[item["category"]])
            for col in range(1, 8):
                sheet.cell(row, col).fill = fill
            for col in range(8, 14):
                sheet.cell(row, col).fill = PatternFill("solid", fgColor="FFFFFF")
            if item["category"] != "场景":
                for col in range(11, 14):
                    sheet.cell(row, col).fill = PatternFill("solid", fgColor="ECEBE8")
            if item["assetType"] == "变体":
                for col in (5, 6):
                    sheet.cell(row, col).fill = PatternFill("solid", fgColor="EED9A7")
                    sheet.cell(row, col).font = Font(name="Microsoft YaHei", bold=True, color="6B3C1F", size=10)
            sheet.row_dimensions[row].height = 90
        elif item["kind"] == "subtotal":
            for col in range(1, 14):
                sheet.cell(row, col).fill = PatternFill("solid", fgColor="E3C996")
                sheet.cell(row, col).font = Font(name="Microsoft YaHei", bold=True, color="4B2A20", size=10)
                sheet.cell(row, col).border = make_border("B99762")
            sheet.row_dimensions[row].height = 28
        else:
            for col in range(1, 14):
                sheet.cell(row, col).fill = PatternFill("solid", fgColor=burgundy)
                sheet.cell(row, col).font = Font(name="Microsoft YaHei", bold=True, color=gold, size=10)
                sheet.cell(row, col).border = make_border(burgundy, "medium")
            sheet.row_dimensions[row].height = 32

        for col in range(1, 14):
            cell = sheet.cell(row, col)
            if item["kind"] == "asset":
                cell.border = make_border()
                if col not in (5, 6) or item["assetType"] != "变体":
                    cell.font = Font(name="Microsoft YaHei", color="2F2A27", size=10)
            cell.alignment = Alignment(
                horizontal="center" if col <= 3 or col == 5 or item["kind"] != "asset" else "left",
                vertical="center",
                wrap_text=True,
            )

    widths = {"A": 9, "B": 9, "C": 9, "D": 32, "E": 12, "F": 32, "G": 48, "H": 20, "I": 20, "J": 20, "K": 20, "L": 20, "M": 20}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "H5"
    sheet.auto_filter.ref = f"A4:M{end_row}"
    table = Table(displayName="AssetInventoryTable", ref=f"A4:M{end_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False, showColumnStripes=False)
    sheet.add_table(table)
    sheet.print_area = f"A1:M{end_row}"
    sheet.print_title_rows = "$1:$4"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()
    return {"rows": end_row, "assetRows": sum(1 for item in meta if item["kind"] == "asset")}


def write_html_preview(data: dict[str, Any], preview_dir: Path) -> Path:
    rows, meta = asset_rows(data)
    preview_dir.mkdir(parents=True, exist_ok=True)
    output = preview_dir / "preview.html"
    body: list[str] = []
    for row, item in zip(rows, meta):
        css_class = f"row-{item['kind']}"
        if item["kind"] == "asset":
            css_class += f" category-{item['category']}"
        cells = "".join(f"<td>{html.escape(str(value or ''))}</td>" for value in row)
        body.append(f'<tr class="{css_class}">{cells}</tr>')
    header_cells = "".join(f"<th>{html.escape(value)}</th>" for value in HEADERS)
    document = f"""<!doctype html><meta charset="utf-8"><title>{html.escape(data['projectTitle'])} 资产统计表预览</title>
<style>
body{{font-family:-apple-system,BlinkMacSystemFont,'Microsoft YaHei',sans-serif;margin:24px;background:#fbf7ef;color:#2f2a27}}
table{{border-collapse:collapse;background:white;min-width:1800px}}th,td{{border:1px solid #d9d1c7;padding:8px;vertical-align:middle;white-space:pre-wrap}}
th{{background:#7a2e2e;color:white;position:sticky;top:0}}td:nth-child(n+8){{min-width:120px;height:72px}}
.category-人物 td:nth-child(-n+7){{background:#f7e8e3}}.category-场景 td:nth-child(-n+7){{background:#f5e9cb}}.category-道具 td:nth-child(-n+7){{background:#e9f0e8}}
.category-人物 td:nth-child(n+11),.category-道具 td:nth-child(n+11){{background:#ecebe8}}.row-subtotal td{{background:#e3c996;font-weight:bold}}.row-grandTotal td{{background:#351719;color:#f1d39b;font-weight:bold}}
</style><h1>{html.escape(data['projectTitle'])} AI真人剧资产统计表</h1><div style="overflow:auto"><table><thead><tr>{header_cells}</tr></thead><tbody>{''.join(body)}</tbody></table></div>"""
    output.write_text(document, encoding="utf-8")
    return output


def inspect_workbook(output: Path, expected_asset_rows: int) -> dict[str, Any]:
    workbook = load_workbook(output, data_only=False)
    issues: list[str] = []
    if workbook.sheetnames != ["资产统计表"]:
        issues.append(f"unexpected sheets: {workbook.sheetnames}")
    sheet = workbook[workbook.sheetnames[0]]
    headers = [sheet.cell(4, col).value for col in range(1, 14)]
    if headers != HEADERS:
        issues.append("header mismatch")
    actual_asset_rows = sum(1 for row in range(5, sheet.max_row + 1) if sheet.cell(row, 3).value in CATEGORY_ORDER)
    if actual_asset_rows != expected_asset_rows:
        issues.append(f"asset row mismatch: expected {expected_asset_rows}, got {actual_asset_rows}")
    for row in range(5, sheet.max_row + 1):
        category = sheet.cell(row, 3).value
        if category in {"小计", "总计"} and not str(sheet.cell(row, 7).value or "").startswith("="):
            issues.append(f"row {row} is missing its total formula")
        for col in range(1, 14):
            value = sheet.cell(row, col).value
            if isinstance(value, str) and re.search(r"#REF!|#DIV/0!|#VALUE!|#NAME\?|#N/A", value):
                issues.append(f"formula error token at {sheet.cell(row, col).coordinate}")
    result = {
        "sheets": workbook.sheetnames,
        "assetRows": actual_asset_rows,
        "maxRow": sheet.max_row,
        "freezePanes": str(sheet.freeze_panes),
        "issues": issues,
        "passed": not issues,
    }
    workbook.close()
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract scripts and build Hermes-compatible AI live-action asset workbooks")
    subparsers = parser.add_subparsers(dest="command", required=True)
    extract_parser = subparsers.add_parser("extract", help="Extract text from DOC, DOCX, PDF, TXT, or Markdown")
    extract_parser.add_argument("source", type=Path)
    extract_parser.add_argument("--output", required=True, type=Path)
    build_parser = subparsers.add_parser("build", help="Build the checked asset workbook from reviewed JSON")
    build_parser.add_argument("--input", required=True, type=Path)
    build_parser.add_argument("--output", required=True, type=Path)
    build_parser.add_argument("--preview-dir", type=Path)
    args = parser.parse_args()

    if args.command == "extract":
        text = extract_script_text(args.source)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(text, encoding="utf-8")
        print(json.dumps({"source": str(args.source), "output": str(args.output), "characters": len(text)}, ensure_ascii=False))
        return

    data = validate_and_normalize(json.loads(args.input.read_text(encoding="utf-8")))
    build_result = build_workbook(data, args.output)
    inspection = inspect_workbook(args.output, build_result["assetRows"])
    preview = str(write_html_preview(data, args.preview_dir)) if args.preview_dir else None
    result = {"output": str(args.output), "preview": preview, **build_result, "inspection": inspection}
    print(json.dumps(result, ensure_ascii=False))
    if not inspection["passed"]:
        raise SystemExit(2)


if __name__ == "__main__":
    main()
