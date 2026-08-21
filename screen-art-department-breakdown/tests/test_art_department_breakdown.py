from __future__ import annotations

import sys
import tempfile
import unittest
from collections import defaultdict
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


SKILL_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL_ROOT / "scripts"))

import art_department_breakdown as builder  # noqa: E402


class SceneHeadingTests(unittest.TestCase):
    def parse(self, heading: str, last_day: str = "日", last_io: str = "内") -> dict | None:
        return builder.parse_heading(heading, 1, defaultdict(int), last_day, last_io, 1)

    def test_marker_first_heading_keeps_day_io_and_location_in_their_columns(self) -> None:
        parsed = self.parse("1-1 日 内 张家客厅")

        self.assertIsNotNone(parsed)
        self.assertEqual(parsed["scene"], "1-1")
        self.assertEqual(parsed["day"], "日")
        self.assertEqual(parsed["inside_outside"], "内")
        self.assertEqual(parsed["location"], "张家客厅")

    def test_location_first_heading_variants_are_recognized(self) -> None:
        headings = [
            "1-1 张家客厅 日 内",
            "1-1 张家客厅/日/内",
            "1-1 张家客厅（日 内）",
            "1-1 张家客厅 日内",
        ]

        for heading in headings:
            with self.subTest(heading=heading):
                self.assertTrue(builder.looks_like_heading(heading))
                parsed = self.parse(heading)
                self.assertIsNotNone(parsed)
                self.assertEqual(parsed["day"], "日")
                self.assertEqual(parsed["inside_outside"], "内")
                self.assertEqual(parsed["location"], "张家客厅")

    def test_heading_without_day_inherits_previous_day_without_shifting_fields(self) -> None:
        parsed = self.parse("1-3 内 村委会会议室", last_day="夜")

        self.assertIsNotNone(parsed)
        self.assertEqual(parsed["day"], "夜")
        self.assertEqual(parsed["inside_outside"], "内")
        self.assertEqual(parsed["location"], "村委会会议室")
        self.assertIn("沿用前场", parsed["warning"])

    def test_prefixed_episode_title_and_nonbreaking_hyphen_are_supported(self) -> None:
        self.assertEqual(builder.detect_episode("《月满西楼》第1集"), 1)

        parsed = self.parse("1‑1 大街上 深夜 外")

        self.assertIsNotNone(parsed)
        self.assertEqual(parsed["scene"], "1-1")
        self.assertEqual(parsed["day"], "夜")
        self.assertEqual(parsed["inside_outside"], "外")
        self.assertEqual(parsed["location"], "大街上")

    def test_midday_and_trailing_heading_notes_are_preserved(self) -> None:
        cases = [
            ("17-7 城区街巷 午后 外", "城区街巷", "日", "外", ""),
            ("3-2 街上 夜 外【闪回】", "街上", "夜", "外", "闪回"),
            ("24-10 沪上城楼 日 外（大结局）", "沪上城楼", "日", "外", "大结局"),
        ]

        for heading, location, day, inside_outside, note in cases:
            with self.subTest(heading=heading):
                self.assertTrue(builder.looks_like_heading(heading))
                parsed = self.parse(heading)
                self.assertIsNotNone(parsed)
                self.assertEqual(parsed["location"], location)
                self.assertEqual(parsed["day"], day)
                self.assertEqual(parsed["inside_outside"], inside_outside)
                if note:
                    self.assertIn(note, parsed["warning"])

    def test_numbered_montage_without_day_or_io_is_kept(self) -> None:
        heading = "4-6 蒙太奇组镜"

        self.assertTrue(builder.looks_like_heading(heading))
        parsed = self.parse(heading)

        self.assertIsNotNone(parsed)
        self.assertEqual(parsed["scene"], "4-6")
        self.assertEqual(parsed["location"], "蒙太奇组镜")
        self.assertEqual(parsed["inside_outside"], "内外")
        self.assertIn("未标日夜内外", parsed["warning"])

    def test_international_screenplay_heading_is_supported(self) -> None:
        heading = "12. INT. LABORATORY - NIGHT"

        self.assertTrue(builder.looks_like_heading(heading))
        parsed = self.parse(heading)

        self.assertIsNotNone(parsed)
        self.assertEqual(parsed["scene"], "1-12")
        self.assertEqual(parsed["location"], "LABORATORY")
        self.assertEqual(parsed["day"], "夜")
        self.assertEqual(parsed["inside_outside"], "内")

    def test_unnumbered_feature_film_heading_uses_document_sequence(self) -> None:
        parsed = builder.parse_heading(
            "INT. OBSERVATION DECK - DAY",
            None,
            defaultdict(int),
            "夜",
            "外",
            7,
        )

        self.assertIsNotNone(parsed)
        self.assertEqual(parsed["scene"], "7")
        self.assertEqual(parsed["location"], "OBSERVATION DECK")
        self.assertIn("无显式场号", parsed["warning"])

    def test_scene_prefixed_feature_film_heading_is_supported(self) -> None:
        parsed = builder.parse_heading(
            "SCENE 12 - EXT. CITY SQUARE - DUSK",
            None,
            defaultdict(int),
            "日",
            "内",
            1,
        )

        self.assertIsNotNone(parsed)
        self.assertEqual(parsed["scene"], "12")
        self.assertEqual(parsed["location"], "CITY SQUARE")
        self.assertEqual(parsed["day"], "夜")
        self.assertEqual(parsed["inside_outside"], "外")

    def test_chinese_numeral_scene_heading_is_supported(self) -> None:
        heading = "第三场 张家客厅 日 内"

        self.assertTrue(builder.looks_like_heading(heading))
        parsed = self.parse(heading)

        self.assertIsNotNone(parsed)
        self.assertEqual(parsed["scene"], "1-3")
        self.assertEqual(parsed["location"], "张家客厅")


class PropExtractionTests(unittest.TestCase):
    def scene(self, people: str, location: str, body: str) -> dict:
        return {
            "location": location,
            "people_raw": people,
            "people": builder.split_people(people),
            "body": [body],
            "body_text": body,
        }

    def test_names_and_location_substrings_do_not_create_props(self) -> None:
        cases = [
            ("人物：钱浩", "张家客厅", "钱浩走进屋内。"),
            ("人物：包强", "办公室", "包强坐下。"),
            ("人物：李明", "生产车间", "李明查看机器。"),
            ("人物：沈月白", "妓院门外", "沈月白带来一条关键消息。"),
            ("人物：沈月白", "妓院门外", "沈月白说：是相熟的人力车夫送她回家。"),
            ("人物：沈月白", "街头路口", "屏幕字幕缓缓浮现：1937年7月，上海。"),
        ]

        for people, location, body in cases:
            with self.subTest(people=people, location=location):
                props = builder.extract_scene_props(self.scene(people, location, body))
                self.assertEqual(props, [])

    def test_source_backed_money_bag_and_vehicle_props_are_still_detected(self) -> None:
        scene = self.scene(
            "人物：钱浩、包强、李明",
            "张家门口",
            "钱浩掏出一沓钱，包强拎起手提包，李明开车离开。",
        )

        names = [prop["name"] for prop in builder.extract_scene_props(scene)]

        self.assertIn("现金/红包/银行卡", names)
        self.assertIn("行李/包袋", names)
        self.assertIn("车辆/车钥匙", names)

    def test_period_drama_props_are_named_and_ammunition_is_not_medical(self) -> None:
        scene = self.scene(
            "人物：沈月白",
            "报社",
            "沈月白举起相机，读完请柬和报纸，又看见桌上的手枪、弹药、炸药与引线。",
        )

        names = [prop["name"] for prop in builder.extract_scene_props(scene)]

        self.assertIn("相机/胶卷", names)
        self.assertIn("信件/纸条/请柬", names)
        self.assertIn("报纸/报刊", names)
        self.assertIn("枪械/弹药", names)
        self.assertIn("炸药/引线", names)
        self.assertNotIn("医药用品", names)

    def test_period_transport_smoking_and_handbag_props_are_detected(self) -> None:
        scene = self.scene(
            "人物：张兰梅",
            "上海街头",
            "张兰梅从人力三轮车上走下，点燃香烟，手包掉落在地。",
        )

        names = [prop["name"] for prop in builder.extract_scene_props(scene)]

        self.assertIn("车辆/车钥匙", names)
        self.assertIn("烟具/香烟", names)
        self.assertIn("行李/包袋", names)

    def test_riding_a_public_vehicle_does_not_assign_it_to_the_passenger(self) -> None:
        scene = self.scene(
            "人物：张兰梅",
            "上海街头",
            "张兰梅从人力三轮车上走下，人力车夫驾车离去。",
        )

        self.assertEqual(builder.scene_prop_text(scene), "车辆/车钥匙")

    def test_props_only_discussed_in_dialogue_are_not_counted_as_visible(self) -> None:
        scene = self.scene(
            "人物：老詹姆斯",
            "沈府花园",
            "老詹姆斯：只要你点头，枪支弹药、药品和轿车都可以合作。",
        )

        self.assertEqual(builder.extract_scene_props(scene), [])


class PeopleInferenceTests(unittest.TestCase):
    def test_action_and_graphic_lines_are_not_treated_as_character_names(self) -> None:
        speakers = builder.dialogue_speakers(
            [
                "沈月白展开纸条，纸上写：松江路花园洋房301。",
                "信上字迹潦草凌乱：爹，我知错了。",
                "头版黑体大字醒目震撼：《三据点全线清缴》。",
                "覃西楼点头：这件事交给我。",
                "众人齐声应和：同心守土！",
            ]
        )

        self.assertEqual(speakers, ["覃西楼", "众人"])

    def test_empty_shot_is_explicitly_marked_as_having_no_people(self) -> None:
        scene = {
            "location": "城市空镜",
            "people_raw": "",
            "people": [],
            "body_text": "城市空镜缓缓升起。",
        }

        self.assertEqual(builder.people_and_groups(scene), "无人物（空镜）")


class SceneGroupingTests(unittest.TestCase):
    def test_common_compound_locations_are_grouped_under_their_practical_main_set(self) -> None:
        cases = [
            ("沈府中堂", "内", ("沈府", "中堂")),
            ("沈府花园内", "外", ("沈府", "花园内")),
            ("徐里洋房门口", "外", ("徐里洋房", "门口")),
            ("徐里洋房内", "内", ("徐里洋房", "徐里洋房内景")),
            ("出租屋门外", "外", ("出租屋", "门外")),
            ("警局审讯室", "内", ("警局", "审讯室")),
            ("警局走廊", "内", ("警局", "走廊")),
            ("监狱探视室", "内", ("监狱", "探视室")),
            ("主控房内", "内", ("主控房", "主控房内景")),
            ("屋内", "内", ("未明确房屋（需确认）", "屋内")),
            ("水厂主控房门口", "外", ("水厂", "主控房门口")),
            ("审讯室门外", "内", ("未明确审讯区（需确认）", "审讯室门外")),
            ("长三堂子，妓女阿妹房间", "内", ("长三堂子", "妓女阿妹房间")),
            ("长三堂子厢房", "内", ("长三堂子", "厢房")),
            ("长三堂子阿妹房间", "内", ("长三堂子", "阿妹房间")),
        ]

        for location, inside_outside, expected in cases:
            with self.subTest(location=location):
                self.assertEqual(builder.split_location_for_scouting(location, inside_outside), expected)

    def test_named_venue_functional_areas_reuse_the_same_main_set(self) -> None:
        cases = [
            ("沈府后门", "外", ("沈府", "后门")),
            ("沈府仓库", "内", ("沈府", "仓库")),
            ("报社印刷间", "内", ("报社", "印刷间")),
            ("警局情报室", "内", ("警局", "情报室")),
            ("水厂巡查道", "外", ("水厂", "巡查道")),
            ("覃家宅院书房", "内", ("覃家", "书房")),
        ]

        for location, inside_outside, expected in cases:
            with self.subTest(location=location):
                self.assertEqual(builder.split_location_for_scouting(location, inside_outside), expected)

        self.assertEqual(
            builder.split_location_for_scouting("码头帮仓库", "内"),
            ("码头帮仓库", "码头帮仓库内景"),
        )

    def test_urban_street_dressing_does_not_use_rural_defaults(self) -> None:
        dressing = builder.basic_set_dressing("上海街头", "街头外景", "上海街头")

        self.assertTrue(any("店招" in item or "城市" in item for item in dressing))
        self.assertFalse(any("村" in item or "农具" in item for item in dressing))

    def test_period_interiors_and_brothel_do_not_receive_modern_or_rural_defaults(self) -> None:
        living_room = builder.basic_set_dressing("沈府", "客厅", "沈府客厅")
        central_hall = builder.basic_set_dressing("沈府", "中堂", "沈府中堂")
        brothel = builder.basic_set_dressing("长三堂子", "妓院外景", "长三堂子妓院外")

        self.assertFalse(any("电视" in item for item in living_room))
        self.assertTrue(any("主位" in item or "屏风" in item for item in central_hall))
        self.assertTrue(any("灯" in item or "招牌" in item for item in brothel))
        self.assertFalse(any("农具" in item for item in brothel))

    def test_generic_offices_and_hospitals_do_not_match_village_or_home_rules(self) -> None:
        office_rule = builder.match_location_rule("警局办公室")
        hospital_rule = builder.match_location_rule("租界医院")

        self.assertIsNotNone(office_rule)
        self.assertIsNotNone(hospital_rule)
        self.assertNotIn("村委会", office_rule["graphic"])
        self.assertIn("医疗", hospital_rule["wall"])

    def test_doorway_and_newsroom_receive_space_appropriate_basic_dressing(self) -> None:
        doorway = builder.basic_set_dressing("出租屋", "门外", "出租屋门外")
        newsroom = builder.basic_set_dressing("报社", "报社内景", "报社")

        self.assertTrue(any("门" in item or "外立面" in item for item in doorway))
        self.assertFalse(any("床" in item for item in doorway))
        self.assertTrue(any("办公" in item or "稿件" in item for item in newsroom))


class PdfExtractionTests(unittest.TestCase):
    def test_stsong_body_text_is_not_discarded_as_a_watermark(self) -> None:
        words = [
            {"text": "1-1", "fontname": "STSong", "size": 12, "non_stroking_color": 0, "top": 10, "x0": 0},
            {"text": "日", "fontname": "STSong", "size": 12, "non_stroking_color": 0, "top": 10, "x0": 30},
            {"text": "内", "fontname": "STSong", "size": 12, "non_stroking_color": 0, "top": 10, "x0": 45},
            {"text": "张家客厅", "fontname": "STSong", "size": 12, "non_stroking_color": 0, "top": 10, "x0": 60},
        ]

        class FakePage:
            def extract_words(self, **_kwargs):
                return words

        class FakePdf:
            pages = [FakePage()]

            def __enter__(self):
                return self

            def __exit__(self, *_args):
                return None

        class FakePdfPlumber:
            @staticmethod
            def open(_path):
                return FakePdf()

        original = builder.pdfplumber
        builder.pdfplumber = FakePdfPlumber()
        try:
            lines = builder.pdf_lines(Path("fixture.pdf"))
        finally:
            builder.pdfplumber = original

        self.assertEqual(lines, [{"text": "1-1日内张家客厅", "page": 1}])


class WorkbookIntegrationTests(unittest.TestCase):
    def test_long_wrapped_cells_expand_their_row_height(self) -> None:
        workbook = builder.Workbook()
        sheet = workbook.active
        sheet.append(["内容"])
        sheet.append(["很长的场景汇总内容" * 20])

        builder.style_sheet(sheet, {1: 20})

        self.assertGreater(sheet.row_dimensions[2].height, 38)

    def test_docx_line_breaks_create_scene_boundaries_and_people_are_inferred(self) -> None:
        document = Document()
        document.add_paragraph(
            "《月满西楼》第1集\n"
            "1‑1 大街上 深夜 外\n"
            "沈月白举起相机。\n"
            "覃西楼（画外）：放下。\n"
            "沈月白：为什么？\n"
            "字幕：1937年6月，上海。\n"
            "大结局字幕：愿山河无恙，人间皆安。\n"
            "1‑2 报社编辑部 午后 内\n"
            "沈月白推门进入报社。"
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "line-breaks.docx"
            document.save(source)
            scenes = builder.parse_script(source)

        self.assertEqual([scene["scene"] for scene in scenes], ["1-1", "1-2"])
        self.assertIn("沈月白", scenes[0]["people"])
        self.assertIn("覃西楼", scenes[0]["people"])
        self.assertNotIn("字幕", scenes[0]["people"])
        self.assertNotIn("大结局字幕", scenes[0]["people"])
        self.assertNotIn("报社编辑部", scenes[0]["body_text"])

    def test_docx_table_content_is_parsed_in_document_order(self) -> None:
        document = Document()
        table = document.add_table(rows=4, cols=1)
        table.cell(0, 0).text = "第一集"
        table.cell(1, 0).text = "1-1 张家客厅 日 内"
        table.cell(2, 0).text = "人物：张三"
        table.cell(3, 0).text = "张三拿起桌上的文件。"

        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "table-script.docx"
            document.save(source)
            scenes = builder.parse_script(source)

        self.assertEqual([scene["scene"] for scene in scenes], ["1-1"])
        self.assertEqual(scenes[0]["location"], "张家客厅")
        self.assertIn("张三拿起桌上的文件", scenes[0]["body_text"])

    def test_unparsed_heading_candidates_are_reported(self) -> None:
        document = Document()
        document.add_paragraph("正文")
        document.add_paragraph("SCENE XII // UNKNOWN FORMAT")
        document.add_paragraph("1-1 张家客厅 日 内")
        document.add_paragraph("人物：张三")

        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "diagnostics.docx"
            document.save(source)
            diagnostics = []
            scenes = builder.parse_script(source, diagnostics)

        self.assertEqual([scene["scene"] for scene in scenes], ["1-1"])
        self.assertEqual(diagnostics, ["疑似场景标题未识别：SCENE XII // UNKNOWN FORMAT"])

    def test_common_headings_flow_into_four_separate_workbooks_correctly(self) -> None:
        document = Document()
        for paragraph in [
            "正文",
            "第一集",
            "1-1 日 内 张家客厅",
            "人物：张三、李四",
            "△张三拿起手机，递给李四一份文件。",
            "1-2 张家门口 夜 外",
            "人物：张三",
            "△张三开车离开。",
            "1-3 村委会会议室 日 内",
            "人物：李四、村民",
            "△李四在会议桌前查看合同。",
        ]:
            document.add_paragraph(paragraph)

        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "fixture.docx"
            output_folder = Path(temp_dir) / "测试项目-美术拆解"
            document.save(source)
            scenes = builder.parse_script(source)
            outputs = builder.build_workbooks(source, scenes, output_folder, "测试项目")

            self.assertEqual(builder.validate_output_folder(output_folder, scenes), [])

            self.assertEqual(
                sorted(path.name for path in output_folder.glob("*.xlsx")),
                ["01-顺场表.xlsx", "02-分场表.xlsx", "03-场景表.xlsx", "04-置景道具执行表.xlsx"],
            )

            expected_sheet_names = {
                "顺场表": "01-顺场表.xlsx",
                "分场表": "02-分场表.xlsx",
                "场景表": "03-场景表.xlsx",
                "置景道具执行表": "04-置景道具执行表.xlsx",
            }
            for sheet_name, filename in expected_sheet_names.items():
                with self.subTest(sheet_name=sheet_name):
                    workbook = load_workbook(output_folder / filename)
                    self.assertEqual(workbook.sheetnames, [sheet_name])
                    self.assertEqual(len(workbook[sheet_name].merged_cells.ranges), 0)
                    workbook.close()

            workbook = load_workbook(outputs["顺场表"])
            sheet = workbook["顺场表"]
            headers = [cell.value for cell in sheet[1]]
            first_row = [sheet.cell(2, column).value for column in range(1, 8)]
            workbook.close()

            execution = load_workbook(outputs["置景道具执行表"])
            execution_headers = [cell.value for cell in execution["置景道具执行表"][1]]
            execution.close()

        self.assertEqual([scene["scene"] for scene in scenes], ["1-1", "1-2", "1-3"])
        self.assertEqual(
            headers,
            [
                "序号",
                "场次",
                "剧本原场景名",
                "归并主场景",
                "功能区/次场景",
                "日/夜",
                "内/外",
                "人物/群演",
                "美术相关剧情/画面要求",
                "场景状态/连续性",
                "陈设道具",
                "戏用道具",
                "特殊/安全项",
                "待确认",
            ],
        )
        self.assertEqual(first_row, [1, "1-1", "张家客厅", "张家", "客厅", "日", "内"])
        self.assertIn("涉及场次", execution_headers)
        self.assertIn("空间处理/改造要求", execution_headers)
        self.assertNotIn("备注", execution_headers)


class GeneralizationTests(unittest.TestCase):
    def test_state_and_continuity_are_inferred_from_generic_visual_changes(self) -> None:
        scene = {
            "location": "客厅",
            "day": "夜",
            "body_text": "房间被砸得一片狼藉，地面满是积水，窗帘正在燃烧。",
            "warning": "",
        }

        state = builder.scene_state_and_continuity(scene)

        self.assertIn("破损/损坏版本", state)
        self.assertIn("潮湿/积水版本", state)
        self.assertIn("燃烧/烟雾版本", state)

    def test_sourcing_uses_production_attributes_instead_of_item_names(self) -> None:
        props = [
            {"name": "陌生图文物", "category": "关键剧情/图文道具"},
            {"name": "陌生专业设备", "category": "专业道具"},
            {"name": "陌生易耗物", "category": "消耗道具"},
        ]

        buckets = builder.sourcing_buckets_for_props(props)

        self.assertEqual(buckets["fabricate"], ["陌生图文物"])
        self.assertEqual(buckets["rent"], ["陌生专业设备"])
        self.assertEqual(buckets["purchase"], ["陌生易耗物"])

    def test_reference_prompts_do_not_assume_a_rural_genre(self) -> None:
        prompts = builder.build_reference_prompts("未来城", "飞船 实验室 人工智能")

        self.assertIn("未来城", prompts)
        self.assertNotIn("乡村", prompts)
        self.assertNotIn("现实主义", prompts)


if __name__ == "__main__":
    unittest.main()
