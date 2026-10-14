from __future__ import annotations

import sys
import json
from copy import deepcopy
import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


SKILL_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL_ROOT / "scripts"))

import shunchangbiao_builder as builder  # noqa: E402


class SceneHeadingTests(unittest.TestCase):
    def test_source_suffixes_survive_word_and_pdf_heading_parsers(self) -> None:
        for heading in ("1-1a 客厅 日 内", "1-1a客厅 日 内", "1-1aa 客厅 日 内", "1-1A 客厅·日·内"):
            with self.subTest(heading=heading):
                parsed = builder.parse_heading(heading, 1, {}, "", "")
                self.assertEqual(parsed[0], "1-1A" if "A" in heading else ("1-1aa" if "aa" in heading else "1-1a"))
                self.assertEqual(parsed[3], "客厅")
                self.assertTrue(builder.is_pdf_heading(heading))
                pdf_parsed = builder.parse_pdf_heading(heading, None, {}, "", "")
                self.assertEqual(pdf_parsed[:4], parsed[:4])

    def test_numbered_empty_shot_is_always_a_scene_boundary(self) -> None:
        parsed = builder.parse_heading("1-1空镜", 1, {}, "", "")

        self.assertIsNotNone(parsed)
        self.assertEqual(parsed[:4], ("1-1", "需确认", "需确认", "空镜"))
        self.assertIn("日夜需确认", parsed[4])
        self.assertIn("内外需确认", parsed[4])

    def test_trailing_flashback_annotation_does_not_hide_scene(self) -> None:
        parsed = builder.parse_heading(
            "2-2 沈家旧宅小院  夜 外（闪回 同3-1）",
            2,
            {},
            "夜",
            "外",
        )

        self.assertIsNotNone(parsed)
        self.assertEqual(parsed[:4], ("2-2", "夜", "外", "沈家旧宅小院"))
        self.assertEqual(parsed[5], "闪回 同3-1")

    def test_chinese_episode_numbers_over_thirty_are_supported(self) -> None:
        self.assertEqual(builder.cn_episode_to_int("三十一"), 31)
        self.assertEqual(builder.cn_episode_to_int("八十"), 80)
        self.assertEqual(builder.cn_episode_to_int("一百零二"), 102)

    def test_docx_scene_bodies_stop_at_incomplete_headings(self) -> None:
        document = Document()
        for paragraph in [
            "正文",
            "第一集",
            "1-1空镜",
            "△城市外景。",
            "1-2 百乐门外空镜+下城区空镜",
            "△百乐门外。",
            "1-3 百乐门舞厅 夜 内",
            "人物：沈晚晴",
            "△沈晚晴走进舞厅。",
            "第二集",
            "2-1 贫民窟暗巷 夜 外",
            "人物：沈晚晴",
            "△沈晚晴拐入暗巷。",
            "2-2 沈家旧宅 夜 内（闪回）",
            "人物：沈父、幼年沈晚晴",
            "△沈父护住幼年沈晚晴。",
        ]:
            document.add_paragraph(paragraph)

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "fixture.docx"
            document.save(path)
            scenes = builder.parse_docx(path)

        self.assertEqual([scene["scene"] for scene in scenes], ["1-1", "1-2", "1-3", "2-1", "2-2"])
        self.assertNotIn("百乐门外", scenes[0]["body_text"])
        self.assertNotIn("沈父", scenes[3]["body_text"])

    def test_docx_table_content_is_read_in_document_order(self) -> None:
        document = Document()
        document.add_paragraph("正文")
        document.add_paragraph("第一集")
        table = document.add_table(rows=4, cols=1)
        table.cell(0, 0).text = "1-1 办公室 日 内"
        table.cell(1, 0).text = "人物：陆锋"
        table.cell(2, 0).text = "△陆锋查看文件。"
        table.cell(3, 0).text = "1-2 走廊 日 内"

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "table-script.docx"
            document.save(path)
            scenes = builder.parse_docx(path)

        self.assertEqual([scene["scene"] for scene in scenes], ["1-1", "1-2"])
        self.assertIn("陆锋查看文件", scenes[0]["body_text"])

    def test_legacy_doc_is_converted_and_parsed(self) -> None:
        if shutil.which("textutil") is None:
            self.skipTest("textutil is unavailable")
        document = Document()
        for paragraph in ["正文", "第一集", "1-1 办公室 日 内", "人物：陆锋", "△陆锋查看文件。"]:
            document.add_paragraph(paragraph)

        with tempfile.TemporaryDirectory() as tmpdir:
            docx_path = Path(tmpdir) / "legacy-source.docx"
            doc_path = Path(tmpdir) / "legacy-source.doc"
            document.save(docx_path)
            subprocess.run(
                ["textutil", "-convert", "doc", "-output", str(doc_path), str(docx_path)],
                check=True,
                capture_output=True,
                text=True,
            )
            scenes = builder.parse_script(doc_path)

        self.assertEqual([scene["scene"] for scene in scenes], ["1-1"])


class SceneSplitTests(unittest.TestCase):
    def setUp(self) -> None:
        # First transition is mid-paragraph; return to A must remain a third row.
        body = ["△林晓在客厅拿信。△林晓走到院子，保安拦住她。", "△林晓回到客厅读信。"]
        self.source = {
            "scene": "1-1", "source": "fixture.docx", "idx": 0,
            "heading": "1-1 林宅客厅 日 内", "location": "林宅客厅",
            "day": "日", "inside_outside": "内", "people": ["林晓", "保安"],
            "people_raw": "人物：林晓、保安", "heading_note": "", "warning": "",
            "body": body, "body_text": "\n".join(body), "body_line_count_draft": 30,
        }
        text = self.source["body_text"]
        self.plan = {
            "reviewed": True, "reason": "林晓由客厅到院子与保安交涉，然后回客厅读信",
            "source_body_sha256": builder.source_body_digest(self.source),
            "parts": [
                {"start": 0, "location": "林宅客厅", "day": "日", "inside_outside": "内"},
                {"start": text.index("△林晓走到院子"), "location": "林宅院子", "day": "日", "inside_outside": "外"},
                {"start": text.index("△林晓回到客厅"), "location": "林宅客厅", "day": "日", "inside_outside": "内"},
            ],
        }
        self.config = {
            "roles": [{"name": "林晓", "mark": "晓"}], "scene_splits": {"1-1": self.plan},
            "continuity_groups": [{"scenes": ["1-1#1", "1-1#2", "1-1#3"], "reviewed": True,
                                   "reason": "同在林宅，林晓拿信出院与保安交涉后立即返回读信"}],
            "scenes": {},
        }
        for i in range(3):
            self.config["scenes"][f"1-1#{i + 1}"] = {
                "reviewed": True, "summary": ["林晓拿起桌上信件准备到院子", "林晓在院子被保安拦住后交涉", "林晓返回客厅坐下仔细读信件"][i],
                "main_location": "林宅", "body_line_count": [3, 4, 2][i], "hint": "",
                "on_camera": ["林晓"], "special_actors": "保安" if i == 1 else "", "groups": "",
                "vo": "", "os": "", "costume": "", "props": "信件" if i != 1 else "",
                "vehicles": "", "notes": "",
            }

    def test_mid_paragraph_split_preserves_every_character_and_return_visit(self) -> None:
        original = deepcopy(self.source)
        scenes, issues = builder.apply_scene_splits([self.source], self.config)
        self.assertEqual(issues, [])
        self.assertEqual([s["scene"] for s in scenes], ["1-1", "1-1a", "1-1"])
        self.assertEqual([s["row_id"] for s in scenes], ["1-1#1", "1-1#2", "1-1#3"])
        self.assertEqual("".join(s["body_text"] for s in scenes), self.source["body_text"])
        self.assertEqual(self.source, original)
        self.assertEqual([s["location"] for s in scenes], ["林宅客厅", "林宅院子", "林宅客厅"])
        self.assertEqual(scenes[0]["source_body_end"], scenes[1]["source_body_start"])
        for child in scenes:
            self.assertEqual(child["source_scene"], "1-1")
            self.assertEqual(child["source_heading"], self.source["heading"])
            self.assertEqual(child["people"], [])
            self.assertIsNone(child["body_line_count_draft"])

    def test_camera_change_and_dialogue_mentions_stay_unsplit_without_reviewed_move(self) -> None:
        self.source["body_text"] = "△切近景，林晓看向窗外。\n林晓：明天我去医院。"
        scenes, issues = builder.apply_scene_splits([self.source], {"scene_splits": {}})
        self.assertEqual(issues, [])
        self.assertEqual([s["scene"] for s in scenes], ["1-1"])

    def test_invalid_boundaries_reject_all_splits(self) -> None:
        for starts in ([1, 10, 20], [0, 20, 10], [0, 10, 10], [0, -1, 20],
                       [0, 10, 999], [False, 10, 20], [0, 1.5, 20], [0, "10", 20]):
            with self.subTest(starts=starts):
                config = deepcopy(self.config)
                for part, start in zip(config["scene_splits"]["1-1"]["parts"], starts):
                    part["start"] = start
                scenes, issues = builder.apply_scene_splits([self.source], config)
                self.assertTrue(issues)
                self.assertEqual([s["scene"] for s in scenes], ["1-1"])

    def test_stale_source_unknown_parent_and_unreviewed_plan_rejected(self) -> None:
        for field, value in (("source_body_sha256", "old"), ("reviewed", False),
                             ("reason", ""), ("parts", []), ("parts", [None, None])):
            with self.subTest(field=field):
                config = deepcopy(self.config)
                config["scene_splits"]["1-1"][field] = value
                self.assertTrue(builder.apply_scene_splits([self.source], config)[1])
        self.assertTrue(builder.apply_scene_splits([self.source], {"scene_splits": {"2-9": self.plan}})[1])
        self.assertTrue(builder.apply_scene_splits([self.source], {"scene_splits": []})[1])

    def test_number_collisions_and_stale_parent_review_rejected(self) -> None:
        duplicate = dict(self.source, scene="1-1a")
        self.assertTrue(builder.apply_scene_splits([self.source, duplicate], self.config)[1])
        self.assertTrue(builder.apply_scene_splits([self.source, self.source], self.config)[1])
        self.config["scenes"]["1-1"] = self.config["scenes"]["1-1#1"]
        self.assertTrue(builder.apply_scene_splits([self.source], self.config)[1])

    def test_existing_suffix_is_preserved_and_not_automatically_split_again(self) -> None:
        source = dict(self.source, scene="1-1a")
        scenes, issues = builder.apply_scene_splits([source], {"scene_splits": {}})
        self.assertEqual(issues, [])
        self.assertEqual(scenes[0]["scene"], "1-1a")
        self.assertTrue(builder.apply_scene_splits([source], {"scene_splits": {"1-1a": self.plan}})[1])
        self.assertEqual([builder.split_suffix(i) for i in (0, 25, 26, 27)], ["a", "z", "aa", "ab"])

    def test_child_location_and_time_fields_are_required_not_inherited(self) -> None:
        for field, value in (("location", ""), ("day", []), ("inside_outside", {})):
            config = deepcopy(self.config)
            config["scene_splits"]["1-1"]["parts"][1][field] = value
            self.assertTrue(builder.apply_scene_splits([self.source], config)[1])

    def test_strict_requires_split_review_even_if_none_and_rejects_missing_children(self) -> None:
        scenes, _ = builder.apply_scene_splits([self.source], self.config)
        roles = builder.resolve_roles(scenes, self.config)
        self.assertEqual(builder.validation_issues(scenes, self.config, roles, strict=True), [])
        del self.config["scenes"]["1-1#2"]
        self.assertTrue(any("1-1#2：未完成逐场复核" in issue for issue in builder.validation_issues(scenes, self.config, roles, strict=True)))
        del self.config["scene_splits"]
        self.assertTrue(any("场内转场复核" in issue for issue in builder.validation_issues(scenes, self.config, roles, strict=True)))

    def test_both_sheets_share_child_data_and_source_provenance(self) -> None:
        scenes, _ = builder.apply_scene_splits([self.source], self.config)
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "split.xlsx"
            builder.build_workbook(Path("fixture.docx"), scenes, output, self.config)
            wb = load_workbook(output)
            schedule, locations = wb["顺场表"], wb["分场表"]
            rows = {row[0]: row for row in schedule.iter_rows(min_row=2, values_only=True)}
            self.assertEqual(rows, {row[0]: row for row in locations.iter_rows(min_row=2, values_only=True)})
            self.assertEqual([row[3] for row in locations.iter_rows(min_row=2, values_only=True)], ["1-1", "1-1", "1-1a"])
            self.assertEqual([rows[i][8] for i in (1, 2, 3)], [0.1, 0.2, 0.1])
            self.assertEqual([rows[i][10] for i in (1, 2, 3)], [None, "保安", None])
            self.assertEqual(len(rows), 3)
            self.assertNotEqual(rows[1][5], rows[3][5])
            self.assertEqual([row[0] for row in locations.iter_rows(min_row=2, values_only=True)], [1, 3, 2])
            self.assertIn("接1-1（顺场1）", rows[2][-1])
            self.assertIn("续1-1（顺场3）", rows[2][-1])
            self.assertNotIn("#", "".join(str(value) for row in rows.values() for value in row))
            for ws in wb:
                self.assertEqual(ws["D2"].border.top.style, "thick")
                self.assertEqual(ws["D4"].border.bottom.style, "thick")
                self.assertIn("原场：1-1", ws["D2"].comment.text)
                self.assertIn("原场1-1场内拆分", ws.cell(2, ws.max_column).value)
            wb.close()

    def test_cross_main_location_split_does_not_qualify_for_one_continuity_box(self) -> None:
        self.config["scenes"]["1-1#2"]["main_location"] = "警局"
        self.plan["parts"][1]["location"] = "警局门口"
        scenes, issues = builder.apply_scene_splits([self.source], self.config)
        self.assertEqual(issues, [])
        groups, issues = builder.resolve_continuity_groups(scenes, self.config)
        self.assertEqual(groups, [])
        self.assertTrue(any("跨主场景" in issue for issue in issues))

    def test_cli_applies_splits_before_strict_validation_and_reports_source_count(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            script = root / "fixture.docx"
            doc = Document()
            doc.add_paragraph(self.source["heading"])
            for paragraph in self.source["body"]:
                doc.add_paragraph(paragraph)
            doc.save(script)
            config_path, report, dump = root / "config.json", root / "report.json", root / "dump.json"
            config_path.write_text(json.dumps(self.config, ensure_ascii=False))
            command = [sys.executable, str(SKILL_ROOT / "scripts/shunchangbiao_builder.py"),
                       str(script), "--out-dir", str(root / "output"), "--config", str(config_path),
                       "--report", str(report), "--dump-json", str(dump), "--strict"]
            result = subprocess.run(command, capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            record = json.loads(report.read_text())[script.name]
            self.assertEqual(record["source_scene_count"], 1)
            self.assertEqual(record["scene_count"], 3)
            self.assertEqual(record["split_map"], {"1-1": ["1-1", "1-1a", "1-1"]})
            self.assertEqual(record["split_row_ids"], {"1-1": ["1-1#1", "1-1#2", "1-1#3"]})
            self.assertTrue(record["delivery_allowed"])
            self.assertEqual("".join(s["body_text"] for s in json.loads(dump.read_text())[script.name]), self.source["body_text"])
            self.plan["source_body_sha256"] = "stale"
            config_path.write_text(json.dumps(self.config, ensure_ascii=False))
            command[command.index("--out-dir") + 1] = str(root / "invalid")
            result = subprocess.run(command, capture_output=True, text=True)
            self.assertEqual(result.returncode, 2)
            self.assertFalse(list((root / "invalid").glob("*.xlsx")))

    def test_new_spaces_get_letters_and_each_return_reuses_its_number(self) -> None:
        source = dict(self.source, scene="2-7", body_text="甲。乙。丙。丁。戊。")
        places = ["林宅客厅", "林宅院子", "林宅客厅", "林宅厨房", "林宅院子"]
        plan = dict(self.plan, source_body_sha256=builder.source_body_digest(source), parts=[
            {"start": i * 2, "location": place, "day": "日", "inside_outside": "外" if place.endswith("院子") else "内"}
            for i, place in enumerate(places)
        ])
        next_scene = dict(self.source, scene="2-8")
        scenes, issues = builder.apply_scene_splits([source, next_scene], {"scene_splits": {"2-7": plan}})
        self.assertEqual(issues, [])
        self.assertEqual([s["scene"] for s in scenes], ["2-7", "2-7a", "2-7", "2-7b", "2-7a", "2-8"])
        self.assertEqual(len({builder.scene_key(s) for s in scenes}), 6)

    def test_aliases_require_reviewed_space_key_and_unknown_spaces_do_not_reuse(self) -> None:
        self.plan["parts"][2]["location"] = "林家客厅"
        self.plan["parts"][0]["space_key"] = "林宅客厅"
        self.plan["parts"][2]["space_key"] = "林宅客厅"
        scenes, issues = builder.apply_scene_splits([self.source], self.config)
        self.assertEqual(issues, [])
        self.assertEqual([s["scene"] for s in scenes], ["1-1", "1-1a", "1-1"])
        self.plan["parts"][2]["space_key"] = "王宅客厅"
        scenes, _ = builder.apply_scene_splits([self.source], self.config)
        self.assertEqual([s["scene"] for s in scenes], ["1-1", "1-1a", "1-1b"])
        for part in self.plan["parts"]:
            part["space_key"] = "需确认"
        scenes, _ = builder.apply_scene_splits([self.source], self.config)
        self.assertEqual([s["scene"] for s in scenes], ["1-1", "1-1a", "1-1b"])
        self.plan["parts"][1]["space_key"] = []
        self.assertTrue(builder.apply_scene_splits([self.source], self.config)[1])

    def test_continuity_targets_exact_occurrence_after_location_sort(self) -> None:
        self.config["continuity_groups"][0]["scenes"] = ["1-1#1", "1-1#2"]
        scenes, _ = builder.apply_scene_splits([self.source], self.config)
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "split.xlsx"
            builder.build_workbook(Path("fixture.docx"), scenes, output, self.config)
            wb = load_workbook(output)
            self.assertEqual(wb["顺场表"]["D4"].border.left.style, "thin")
            self.assertEqual(wb["分场表"]["D3"].border.left.style, "thin")
            for cell in (wb["分场表"]["D2"], wb["分场表"]["D4"]):
                self.assertEqual(cell.border.top.style, "thick")
                self.assertEqual(cell.border.bottom.style, "thick")
            wb.close()

    def test_unreviewed_duplicate_labels_and_duplicate_internal_keys_still_fail(self) -> None:
        issues = builder.validation_issues([self.source, dict(self.source)], {}, [], strict=False)
        self.assertTrue(any("重复" in issue for issue in issues))
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "invalid.xlsx"
            with self.assertRaisesRegex(ValueError, "唯一行标识"):
                builder.build_workbook(Path("fixture.docx"), [self.source, dict(self.source)], output, {})
            self.assertFalse(output.exists())
        scenes, _ = builder.apply_scene_splits([self.source], self.config)
        scenes[2]["row_id"] = scenes[0]["row_id"]
        self.assertTrue(any("标识重复" in issue for issue in builder.validation_issues(scenes, self.config, [], strict=False)))
        scenes[2]["row_id"] = "1-1#3"
        scenes[2]["split_space_key"] = ["王宅客厅", "内"]
        self.assertTrue(any("不属于" in issue for issue in builder.validation_issues(scenes, self.config, [], strict=False)))


class PeopleAndDepartmentTests(unittest.TestCase):
    def test_group_roles_and_voice_roles_do_not_become_actor_columns(self) -> None:
        self.assertTrue(builder.is_group_name("百乐门舞女"))
        self.assertTrue(builder.is_group_name("治安局队员众"))
        self.assertTrue(builder.is_group_name("厂区稽查队员"))
        self.assertTrue(builder.is_voice_name("巡逻员vo"))

        scenes = [
            {
                "people": ["沈晚晴", "百乐门舞女", "治安局队员众", "巡逻员vo"],
            }
        ]
        self.assertEqual(builder.infer_roles(scenes), [("沈晚晴", "晴")])

    def test_age_variant_prevents_duplicate_base_role_in_flashback(self) -> None:
        scenes = [
            {
                "people": ["沈父", "幼年沈晚晴"],
                "body": ["沈晚晴（虚弱）：爸爸！"],
                "heading_note": "闪回",
            },
            {
                "people": ["沈晚晴"],
                "body": ["△沈晚晴走进舞厅。"],
                "heading_note": "",
            },
        ]

        builder.augment_visible_people(scenes)

        self.assertEqual(scenes[0]["people"], ["沈父", "幼年沈晚晴"])
        self.assertEqual(builder.default_mark("沈晚晴"), "晴")
        self.assertEqual(builder.default_mark("幼年沈晚晴"), "幼")

    def test_costume_hints_are_attributed_to_roles(self) -> None:
        scene = {
            "people_raw": "人物：沈晚晴、陆锋",
            "people": ["沈晚晴", "陆锋"],
            "body": [
                "△沈晚晴右腿接驳处灼伤，额头冷汗。",
                "△陆锋身穿治安局制服。",
            ],
            "body_text": "△沈晚晴右腿接驳处灼伤，额头冷汗。\n△陆锋身穿治安局制服。",
        }

        hints = builder.collect_costume(scene)

        self.assertIn("沈晚晴：", hints)
        self.assertIn("陆锋：制服", hints)
        self.assertIn("灼伤", hints)

    def test_props_exclude_background_furniture_and_keep_story_props(self) -> None:
        scene = {
            "body": ["△房间里摆着椅子，沈晚晴收起口红，拿起酒杯，又将沈家纹路金属残片贴在全息光屏上。"],
            "body_text": "△房间里摆着椅子，沈晚晴收起口红，拿起酒杯，又将沈家纹路金属残片贴在全息光屏上。",
        }

        props = builder.collect_props(scene)

        self.assertIn("沈家纹路金属残片", props)
        self.assertIn("全息光屏画面", props)
        self.assertIn("口红", props)
        self.assertIn("酒杯", props)
        self.assertNotIn("椅子", props)

    def test_groups_are_detected_from_action_text_without_people_line(self) -> None:
        scene = {
            "location": "沈家旧宅小院",
            "people_raw": "",
            "people": [],
            "body_text": "△大批机械武装队员从院外破门突入，瞬间封锁小院。",
        }

        self.assertIn("机械武装队员", builder.collect_groups(scene))

    def test_draft_voice_columns_preserve_explicit_source_types(self) -> None:
        scene = {
            "location": "办公室",
            "people": ["巡逻员vo"],
            "body_text": "陆锋os：这件事没有那么简单。",
            "warning": "",
            "heading_note": "",
        }

        voices = builder.collect_voices(scene)

        self.assertEqual(voices, {"vo": "巡逻员", "os": "陆锋"})
        self.assertNotIn("巡逻员VO", builder.collect_notes(scene))


class DeliveryValidationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.scene = {
            "scene": "1-1",
            "location": "城市空镜",
            "day": "需确认",
            "inside_outside": "需确认",
            "heading_note": "",
            "warning": "日夜需确认；内外需确认",
            "people_raw": "",
            "people": [],
            "body": ["△城市空镜交代下城区环境。"],
            "body_text": "△城市空镜交代下城区环境。",
        }

    def test_strict_delivery_requires_every_scene_to_be_reviewed(self) -> None:
        issues = builder.validation_issues([self.scene], {}, [], strict=True)

        self.assertTrue(any("未完成逐场复核" in issue for issue in issues))

    def test_reviewed_scene_with_explicit_unknowns_can_pass_strict_delivery(self) -> None:
        config = {
            "roles": [],
            "scene_splits": {}, "continuity_groups": [],
            "scenes": {
                "1-1": {
                    "reviewed": True,
                    "summary": "城市空镜交代下城区夜间环境",
                    "main_location": "城市",
                    "body_line_count": 1,
                    "hint": "",
                    "on_camera": [],
                    "vo": "", "os": "", "vehicles": "",
                    "special_actors": "",
                    "groups": "",
                    "costume": "",
                    "props": "",
                    "notes": "日夜、内外需确认",
                }
            }
        }

        self.assertEqual(builder.validation_issues([self.scene], config, [], strict=True), [])

    def test_strict_delivery_requires_all_reviewed_fields_even_when_blank(self) -> None:
        config = {
            "scenes": {
                "1-1": {
                    "reviewed": True,
                    "summary": "城市空镜交代下城区夜间环境",
                    "main_location": "城市",
                    "notes": "日夜、内外需确认",
                }
            }
        }

        issues = builder.validation_issues([self.scene], config, [], strict=True)

        self.assertTrue(any("复核字段不完整" in issue for issue in issues))

    def test_duplicate_role_marks_are_rejected(self) -> None:
        roles = [("沈晚晴", "晴"), ("顾晚晴", "晴")]

        issues = builder.validation_issues([self.scene], {}, roles, strict=False)

        self.assertTrue(any("角色简称重复" in issue for issue in issues))

    def test_reviewed_fields_feed_the_workbook_headers_with_special_actors(self) -> None:
        config = {
            "roles": [{"name": "沈晚晴", "mark": "晴"}],
            "scenes": {
                "1-1": {
                    "reviewed": True,
                    "summary": "城市空镜交代下城区夜间环境",
                    "main_location": "城市",
                    "body_line_count": 1,
                    "hint": "",
                    "on_camera": [],
                    "vo": "", "os": "", "vehicles": "",
                    "special_actors": "",
                    "groups": "流民",
                    "costume": "流民：破旧衣服、脏污妆",
                    "props": "雾气特效设备",
                    "notes": "空镜；日夜、内外需确认；需要分别拍摄城市全景、百乐门外立面和下城区街道空镜",
                }
            },
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "schedule.xlsx"
            builder.build_workbook(Path("fixture.docx"), [self.scene], output, config)
            workbook = load_workbook(output)
            sheet = workbook["顺场表"]
            headers = [cell.value for cell in sheet[1]]
            values = [cell.value for cell in sheet[2]]
            freeze_panes = sheet.freeze_panes
            print_title_rows = sheet.print_title_rows
            row_height = sheet.row_dimensions[2].height
            workbook.close()

        self.assertEqual(
            headers,
            [
                "拍摄顺序",
                "实际拍摄场地",
                "提示",
                "场次",
                "剧本中场景",
                "拍摄内容",
                "日\n/\n夜",
                "内\n/\n外",
                "页\n数",
                "沈\n晚\n晴",
                "特约演员",
                "群演",
                "VO",
                "OS",
                "梳化服提示",
                "道具提示",
                "戏用车",
                "备注/特殊道具/时间",
            ],
        )
        self.assertIsNone(values[2])
        self.assertEqual(values[3], "1-1")
        self.assertEqual(values[5], "城市空镜交代下城区夜间环境")
        self.assertEqual([values[headers.index(name)] for name in ["群演", "梳化服提示", "道具提示"]], ["流民", "流民：破旧衣服、脏污妆", "雾气特效设备"])
        self.assertIn("需要分别拍摄城市全景", values[-1])
        self.assertEqual(freeze_panes, "J2")
        self.assertEqual(print_title_rows, "$1:$1")
        self.assertGreater(row_height, 36)


class SpecialActorTests(unittest.TestCase):
    def setUp(self) -> None:
        self.scenes = []
        # Explicit user-supplied roster: other individual roles stay special.
        self.config = {"roles": [{"name": "沈晚晴", "mark": "晴"}], "scene_splits": {}, "continuity_groups": [], "scenes": {}}
        for number in range(1, 4):
            scene_id = f"1-{number}"
            body = [
                "△远处围观群众面目模糊，陈伯单独拦住沈晚晴。",
                "保安：请先登记。",
                "护士：我带你进去。",
            ]
            self.scenes.append({
                "scene": scene_id, "location": "医院走廊", "day": "日", "inside_outside": "内",
                "people": ["沈晚晴", "陈伯", "保安", "护士", "围观群众"],
                "people_raw": "人物：沈晚晴、陈伯、保安、护士、围观群众",
                "body": body, "body_text": "\n".join(body), "heading_note": "", "warning": "",
            })
            self.config["scenes"][scene_id] = {
                "reviewed": True, "summary": "沈晚晴被陈伯拦住后接受保安询问",
                "main_location": "医院",
                "body_line_count": 1,
                "hint": "",
                "on_camera": ["沈晚晴"], "vo": "", "os": "", "vehicles": "",
                "special_actors": "陈伯、保安、护士", "groups": "围观群众",
                "costume": "", "props": "", "notes": "",
            }

    def test_individual_roles_stay_special_across_multiple_scenes(self) -> None:
        roles = builder.resolve_roles(self.scenes, self.config)
        self.assertEqual(builder.validation_issues(self.scenes, self.config, roles, strict=True), [])
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "schedule.xlsx"
            builder.build_workbook(Path("fixture.docx"), self.scenes, path, self.config)
            workbook = load_workbook(path)
            sheet = workbook["顺场表"]
            headers = [cell.value for cell in sheet[1]]
            special_col = headers.index("特约演员") + 1
            group_col = headers.index("群演") + 1
            for row in range(2, 5):
                self.assertEqual(sheet.cell(row, special_col).value, "陈伯、保安、护士")
                self.assertEqual(sheet.cell(row, group_col).value, "围观群众")
                self.assertEqual(sheet.cell(row, 10).value, "晴")
            self.assertEqual(special_col, 11)
            self.assertTrue(sheet.cell(1, special_col).alignment.wrap_text)
            self.assertEqual(sheet.auto_filter.ref, "A1:R4")
            workbook.close()

    def test_reviewed_specials_are_excluded_from_inferred_role_columns(self) -> None:
        self.config.pop("roles")
        self.assertEqual(builder.resolve_roles(self.scenes, self.config), [("沈晚晴", "晴")])

    def test_explicit_empty_main_roles_are_respected(self) -> None:
        self.config["roles"] = []
        self.assertEqual(builder.resolve_roles(self.scenes, self.config), [])

    def test_strict_build_requires_reviewed_special_actor_field(self) -> None:
        del self.config["scenes"]["1-1"]["special_actors"]
        issues = builder.validation_issues(self.scenes, self.config, [("沈晚晴", "晴")], strict=True)
        self.assertTrue(any("special_actors" in issue for issue in issues))

    def test_strict_build_requires_source_reviewed_role_roster(self) -> None:
        del self.config["roles"]
        issues = builder.validation_issues(self.scenes, self.config, [("沈晚晴", "晴")], strict=True)
        self.assertTrue(any("roles" in issue for issue in issues))

    def test_specials_cannot_also_occupy_main_role_columns(self) -> None:
        self.config["roles"].append({"name": "陈伯", "mark": "伯"})
        roles = builder.resolve_roles(self.scenes, self.config)
        issues = builder.validation_issues(self.scenes, self.config, roles, strict=True)
        self.assertTrue(any("陈伯" in issue and "独立角色列" in issue for issue in issues))

    def test_special_cannot_revert_to_a_group_in_later_scene(self) -> None:
        self.config["scenes"]["1-2"]["special_actors"] = "保安、护士"
        self.config["scenes"]["1-2"]["groups"] = "陈伯、围观群众"
        issues = builder.validation_issues(self.scenes, self.config, [("沈晚晴", "晴")], strict=True)
        self.assertTrue(any("1-2" in issue and "陈伯" in issue and "群演" in issue for issue in issues))

    def test_voice_only_role_does_not_enter_special_actor_column(self) -> None:
        self.config["scenes"]["1-1"]["special_actors"] = "陈伯、保安、护士、医生VO"
        issues = builder.validation_issues(self.scenes, self.config, [("沈晚晴", "晴")], strict=True)
        self.assertTrue(any("医生VO" in issue for issue in issues))


class VoiceAndVehicleTests(unittest.TestCase):
    def setUp(self) -> None:
        self.scene = {
            "scene": "1-1", "location": "院子", "day": "日", "inside_outside": "外",
            "people": ["沈晚晴", "陆锋", "医生（OS）"], "people_raw": "",
            "body": ["沈晚晴（VO）：当年我为什么没有离开？", "陆锋OS（门外）：开门！"],
            "body_text": "沈晚晴（VO）：当年我为什么没有离开？\n陆锋OS（门外）：开门！",
            "heading_note": "", "warning": "",
        }
        self.review = {
            "reviewed": True, "summary": "沈晚晴听见陆锋喊话后走到门口",
            "main_location": "宅院",
            "body_line_count": 1,
            "hint": "",
            "on_camera": ["沈晚晴"], "special_actors": "", "groups": "",
            "vo": "沈晚晴（内心独白）", "os": "陆锋（门外喊话）；医生（镜外说话）",
            "costume": "", "props": "车钥匙", "vehicles": "轿车1辆（院外停放）", "notes": "",
        }
        self.config = {
            "roles": [{"name": "沈晚晴", "mark": "晴"}, {"name": "陆锋", "mark": "锋"}],
            "scene_splits": {}, "continuity_groups": [], "scenes": {"1-1": self.review},
        }

    def test_voice_only_actor_is_not_marked_visible_but_visible_narrator_is(self) -> None:
        roles = builder.resolve_roles([self.scene], self.config)
        self.assertEqual(builder.validation_issues([self.scene], self.config, roles, strict=True), [])
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "voices.xlsx"
            builder.build_workbook(Path("fixture.docx"), [self.scene], path, self.config)
            workbook = load_workbook(path)
            sheet = workbook["顺场表"]
            headers = [cell.value for cell in sheet[1]]
            values = dict(zip(headers, [cell.value for cell in sheet[2]]))
            self.assertEqual(values["沈\n晚\n晴"], "晴")
            self.assertIsNone(values["陆\n锋"])
            self.assertEqual(values["VO"], self.review["vo"])
            self.assertEqual(values["OS"], self.review["os"])
            self.assertEqual(values["戏用车"], self.review["vehicles"])
            self.assertEqual(values["道具提示"], "车钥匙")
            self.assertEqual(headers[-8:], ["特约演员", "群演", "VO", "OS", "梳化服提示", "道具提示", "戏用车", "备注/特殊道具/时间"])
            workbook.close()

    def test_parenthesized_voice_tags_are_not_lost_from_people_list(self) -> None:
        people = builder.split_people("人物：沈晚晴（VO）、陆锋（OS）")
        self.assertTrue(all(builder.is_voice_name(name) for name in people))
        self.assertEqual(builder.infer_roles([{"people": people}]), [])
        self.assertEqual(builder.collect_voices(self.scene), {"vo": "沈晚晴", "os": "医生、陆锋"})

    def test_generic_phone_or_broadcast_sound_is_not_automatically_vo(self) -> None:
        scene = dict(self.scene, people=["医生电话声", "播音员广播声"], body_text="医生电话声：明天见。\n播音员广播声：现在开始。")
        self.assertEqual(builder.collect_voices(scene), {"vo": "", "os": ""})
        notes = builder.collect_notes(scene)
        self.assertIn("需确认", notes)
        self.assertNotIn("含VO", notes)

    def test_reviewed_voice_classification_can_correct_source_with_a_note(self) -> None:
        self.review.update(vo="陆锋（内心独白）", os="", notes="原文标OS，但正文为内心独白，按VO暂列，待编剧确认")
        self.assertEqual(builder.validation_issues([self.scene], self.config, builder.resolve_roles([self.scene], self.config), strict=True), [])

    def test_strict_review_requires_voice_vehicle_and_visible_cast_fields(self) -> None:
        for field in ["hint", "vo", "os", "vehicles", "on_camera"]:
            with self.subTest(field=field):
                saved = self.review.pop(field)
                issues = builder.validation_issues([self.scene], self.config, builder.resolve_roles([self.scene], self.config), strict=True)
                self.assertTrue(any(field in issue for issue in issues))
                self.review[field] = saved

    def test_visible_cast_override_requires_a_list_of_reviewed_individual_roles(self) -> None:
        for value in ["沈晚晴", ["沈晚晴VO"], ["不在主要角色表中的人"]]:
            with self.subTest(value=value):
                self.review["on_camera"] = value
                issues = builder.validation_issues([self.scene], self.config, builder.resolve_roles([self.scene], self.config), strict=True)
                self.assertTrue(any("on_camera" in issue for issue in issues))


class PageCountTests(unittest.TestCase):
    def test_every_three_body_lines_round_up_to_one_tenth(self) -> None:
        expected = {1: 0.1, 2: 0.1, 3: 0.1, 4: 0.2, 6: 0.2, 7: 0.3,
                    27: 0.9, 28: 1.0, 30: 1.0, 31: 1.1, 33: 1.1, 34: 1.2,
                    60: 2.0, 61: 2.1, 90: 3.0}
        for line_count, pages in expected.items():
            with self.subTest(line_count=line_count):
                self.assertEqual(builder.pages_from_body_lines(line_count), pages)

    def test_unknown_or_invalid_count_does_not_invent_pages(self) -> None:
        for value in [None, 0, -3, 3.5, "30", True]:
            with self.subTest(value=value):
                self.assertEqual(builder.pages_from_body_lines(value), "需确认")

    def test_word_paragraph_count_is_not_used_as_a_layout_line_count(self) -> None:
        document = Document()
        document.add_paragraph("第一集")
        document.add_paragraph("1-1 办公室 日 内")
        document.add_paragraph("人物：沈晚晴")
        document.add_paragraph("△" + "沈晚晴翻阅文件。" * 50)
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "wrapped.docx"
            document.save(path)
            scenes = builder.parse_docx(path)
        self.assertIsNone(scenes[0].get("body_line_count_draft"))
        self.assertIn("原排版", scenes[0].get("body_line_count_basis", ""))

    def test_reviewed_count_drives_numeric_excel_value_and_one_decimal_display(self) -> None:
        fixture = VoiceAndVehicleTests()
        fixture.setUp()
        fixture.review["body_line_count"] = 30
        fixture.scene["body_line_count_draft"] = 4
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "pages.xlsx"
            builder.build_workbook(Path("fixture.docx"), [fixture.scene], path, fixture.config)
            workbook = load_workbook(path)
            sheet = workbook["顺场表"]
            page_col = next(cell.column for cell in sheet[1] if str(cell.value).replace("\n", "") == "页数")
            self.assertEqual(sheet.cell(2, page_col).value, 1.0)
            self.assertEqual(sheet.cell(2, page_col).data_type, "n")
            self.assertEqual(sheet.cell(2, page_col).number_format, "0.0")
            workbook.close()

    def test_strict_build_requires_reviewed_count_and_explains_unknown_pages(self) -> None:
        fixture = VoiceAndVehicleTests()
        fixture.setUp()
        roles = builder.resolve_roles([fixture.scene], fixture.config)
        del fixture.review["body_line_count"]
        self.assertTrue(any("body_line_count" in issue for issue in builder.validation_issues([fixture.scene], fixture.config, roles, strict=True)))
        fixture.review["body_line_count"] = None
        self.assertTrue(any("页数需确认" in issue for issue in builder.validation_issues([fixture.scene], fixture.config, roles, strict=True)))
        fixture.review["notes"] = "原排版暂无法核行，页数需确认"
        self.assertEqual(builder.validation_issues([fixture.scene], fixture.config, roles, strict=True), [])
        fixture.review["body_line_count"] = -1
        self.assertTrue(any("body_line_count" in issue for issue in builder.validation_issues([fixture.scene], fixture.config, roles, strict=True)))


class ContinuityBorderTests(unittest.TestCase):
    def setUp(self) -> None:
        self.scenes = []
        self.config = {"roles": [], "scene_splits": {}, "continuity_groups": [], "scenes": {}}
        for number in range(1, 7):
            scene_id = f"1-{number}"
            self.scenes.append({
                "scene": scene_id, "location": "院子", "day": "日", "inside_outside": "外",
                "people": [], "people_raw": "", "body": [], "body_text": "",
                "heading_note": "", "warning": "",
            })
            self.config["scenes"][scene_id] = {
                "reviewed": True, "summary": "沈晚晴跑出院子后陆锋紧随追出",
                "main_location": "同一宅院",
                "body_line_count": 1,
                "hint": "",
                "on_camera": [], "vo": "", "os": "", "vehicles": "",
                "special_actors": "", "groups": "", "costume": "", "props": "", "notes": "原有备注",
            }

    def group(self, *scene_ids: str) -> dict:
        return {"scenes": list(scene_ids), "reviewed": True, "reason": "同一宅院内前场人物跑出，后场紧接追赶，剧情时间及行动连续"}

    def workbook(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "continuity.xlsx"
            builder.build_workbook(Path("fixture.docx"), self.scenes, path, self.config)
            return load_workbook(path)

    def edges(self, cell):
        return tuple(getattr(cell.border, side).style for side in ("left", "right", "top", "bottom"))

    def test_contiguous_group_gets_one_outer_box_only_in_scene_column(self) -> None:
        self.config["continuity_groups"] = [self.group("1-1", "1-2", "1-3")]
        self.config["scenes"]["1-1"]["hint"] = "雨／烟机／武戏"
        self.assertEqual(builder.validation_issues(self.scenes, self.config, [], strict=True), [])
        workbook = self.workbook()
        sheet = workbook["顺场表"]
        self.assertEqual(self.edges(sheet["D2"]), ("thick", "thick", "thick", "thin"))
        self.assertEqual(self.edges(sheet["D3"]), ("thick", "thick", "thin", "thin"))
        self.assertEqual(self.edges(sheet["D4"]), ("thick", "thick", "thin", "thick"))
        for row in sheet:
            for cell in row:
                if cell.coordinate in {"D2", "D3", "D4"}:
                    self.assertEqual(cell.border.left.color.rgb[-6:], "000000")
                    self.assertFalse(cell.font.bold)
                else:
                    self.assertEqual(self.edges(cell), ("thin",) * 4)
        self.assertEqual([sheet.cell(row, 4).value for row in range(2, 8)], [s["scene"] for s in self.scenes])
        self.assertFalse(sheet.merged_cells.ranges)
        self.assertEqual(sheet.max_column, 17)
        self.assertEqual(sheet.freeze_panes, "J2")
        self.assertEqual(sheet["C2"].value, "雨／烟机／武戏")
        self.assertEqual(sheet["Q2"].value, "原有备注；连场1；续1-2")
        workbook.close()

    def test_intercut_scene_is_excluded_and_separate_runs_are_linked(self) -> None:
        self.config["continuity_groups"] = [self.group("1-1", "1-3", "1-4")]
        workbook = self.workbook()
        sheet = workbook["顺场表"]
        self.assertEqual(self.edges(sheet["D2"]), ("thick",) * 4)
        self.assertEqual(self.edges(sheet["D3"]), ("thin",) * 4)
        self.assertEqual(self.edges(sheet["D4"]), ("thick", "thick", "thick", "thin"))
        self.assertEqual(self.edges(sheet["D5"]), ("thick", "thick", "thin", "thick"))
        self.assertEqual(sheet["Q2"].value, "原有备注；连场1；续1-3")
        self.assertEqual(sheet["Q3"].value, "原有备注")
        self.assertEqual(sheet["Q4"].value, "原有备注；连场1；接1-1；续1-4")
        workbook.close()

    def test_touching_groups_keep_their_separate_top_and_bottom_edges(self) -> None:
        self.config["continuity_groups"] = [self.group("1-1", "1-2"), self.group("1-3", "1-4")]
        workbook = self.workbook()
        sheet = workbook["顺场表"]
        self.assertEqual(sheet["D3"].border.bottom.style, "thick")
        self.assertEqual(sheet["D4"].border.top.style, "thick")
        workbook.close()

    def test_same_location_day_and_adjacent_numbers_do_not_auto_group(self) -> None:
        workbook = self.workbook()
        sheet = workbook["顺场表"]
        for row in range(2, 8):
            self.assertEqual(self.edges(sheet.cell(row, 4)), ("thin",) * 4)
        workbook.close()

    def test_strict_build_requires_explicit_continuity_review_even_when_none(self) -> None:
        del self.config["continuity_groups"]
        issues = builder.validation_issues(self.scenes, self.config, [], strict=True)
        self.assertTrue(any("连场" in issue for issue in issues))

    def test_same_main_location_allows_continuous_different_subspaces(self) -> None:
        self.scenes[0]["location"] = "沈宅·客厅"
        self.scenes[1]["location"] = "沈宅·走廊"
        for scene_id in ("1-1", "1-2"):
            self.config["scenes"][scene_id]["main_location"] = "沈宅"
        self.config["continuity_groups"] = [self.group("1-1", "1-2")]
        self.assertEqual(builder.resolve_continuity_groups(self.scenes, self.config), ([["1-1", "1-2"]], []))

    def test_cross_main_location_is_rejected_and_not_boxed(self) -> None:
        self.config["scenes"]["1-1"]["main_location"] = "沈宅"
        self.config["scenes"]["1-2"]["main_location"] = "警局"
        self.config["continuity_groups"] = [self.group("1-1", "1-2")]
        groups, issues = builder.resolve_continuity_groups(self.scenes, self.config)
        self.assertEqual(groups, [])
        self.assertTrue(any("跨主场景" in issue for issue in issues))
        workbook = self.workbook()
        for address in ("D2", "D3"):
            self.assertEqual(self.edges(workbook["顺场表"][address]), ("thin",) * 4)
        workbook.close()

    def test_missing_main_location_fails_continuity_review(self) -> None:
        del self.config["scenes"]["1-2"]["main_location"]
        self.config["continuity_groups"] = [self.group("1-1", "1-2")]
        issues = builder.validation_issues(self.scenes, self.config, [], strict=True)
        self.assertTrue(any("main_location" in issue and "1-2" in issue for issue in issues))

    def test_invalid_groups_are_reported(self) -> None:
        cases = {
            "unknown": [self.group("1-1", "9-9")],
            "duplicate": [self.group("1-1", "1-1")],
            "single": [self.group("1-1")],
            "overlap": [self.group("1-1", "1-2"), self.group("1-2", "1-3")],
            "reversed": [self.group("1-2", "1-1")],
            "unreviewed": [dict(self.group("1-1", "1-2"), reviewed=False)],
            "no_evidence": [dict(self.group("1-1", "1-2"), reason="")],
            "wrong_type": "1-1、1-2",
        }
        for label, groups in cases.items():
            with self.subTest(label=label):
                self.config["continuity_groups"] = groups
                issues = builder.validation_issues(self.scenes, self.config, [], strict=True)
                self.assertTrue(any("连场" in issue for issue in issues))

    def test_unreviewed_group_does_not_get_a_confirmed_border_in_draft(self) -> None:
        self.config["continuity_groups"] = [dict(self.group("1-1", "1-2"), reviewed=False)]
        workbook = self.workbook()
        self.assertEqual(self.edges(workbook["顺场表"]["D2"]), ("thin",) * 4)
        workbook.close()


class LocationSheetTests(unittest.TestCase):
    def setUp(self) -> None:
        fixture = ContinuityBorderTests()
        fixture.setUp()
        self.scenes, self.config = fixture.scenes, fixture.config
        spaces = [
            ("林宅", "林宅·客厅", "内"), ("警局", "警局·大厅", "内"),
            ("林宅", "林宅·走廊", "内"), ("林宅", "林家·客厅", "内"),
            ("林宅", "林宅·客厅", "外"), ("警局", "警局·大厅", "内"),
        ]
        for scene, (main, location, inside) in zip(self.scenes, spaces):
            scene.update(location=location, inside_outside=inside)
            self.config["scenes"][scene["scene"]]["main_location"] = main
        self.config["scenes"]["1-4"]["location_group"] = "林宅·客厅"
        self.scenes[3]["day"] = "夜"

    def workbook(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "two-sheets.xlsx"
            builder.build_workbook(Path("fixture.docx"), self.scenes, path, self.config)
            return load_workbook(path)

    def test_two_views_preserve_every_cell_and_group_by_reviewed_location(self) -> None:
        book = self.workbook()
        self.assertEqual(book.sheetnames, ["顺场表", "分场表"])
        sequential, grouped = book.worksheets
        expected = ["1-1", "1-4", "1-3", "1-5", "1-2", "1-6"]
        self.assertEqual([grouped.cell(row, 4).value for row in range(2, 8)], expected)
        self.assertEqual([grouped.cell(row, 1).value for row in range(2, 8)], [1, 4, 3, 5, 2, 6])
        self.assertEqual([cell.value for cell in sequential[1]], [cell.value for cell in grouped[1]])
        data = {row[3]: row for row in sequential.iter_rows(min_row=2, values_only=True)}
        for row in grouped.iter_rows(min_row=2, values_only=True):
            self.assertEqual(row, data[row[3]])
        self.assertEqual(grouped["E3"].value, "林家·客厅")
        for sheet in book:
            self.assertEqual(sheet.freeze_panes, "J2")
            self.assertEqual(sheet.auto_filter.ref, "A1:Q7")
            self.assertEqual(sheet.print_title_rows, "$1:$1")
            self.assertEqual(sheet.page_setup.orientation, "landscape")
            self.assertFalse(sheet.merged_cells.ranges)
            self.assertTrue(all(sheet.cell(r, 2).value is None for r in range(2, 8)))
        book.close()

    def test_reordered_members_receive_correct_outer_border_and_original_links(self) -> None:
        self.config["continuity_groups"] = [{"scenes": ["1-1", "1-3", "1-4"], "reviewed": True, "reason": "林宅内同一追赶事件从客厅经走廊返回客厅"}]
        book = self.workbook()
        sheet = book["分场表"]
        self.assertEqual(sheet["D2"].border.top.style, "thick")
        self.assertEqual(sheet["D3"].border.top.style, "thin")
        self.assertEqual(sheet["D3"].border.bottom.style, "thin")
        self.assertEqual(sheet["D4"].border.bottom.style, "thick")
        self.assertIn("接1-3", sheet["Q3"].value)
        self.assertIn("接1-1；续1-4", sheet["Q4"].value)
        self.assertEqual(sheet["D5"].border.left.style, "thin")
        book.close()

    def test_same_space_sort_does_not_box_an_unrelated_scene_between_members(self) -> None:
        self.config["continuity_groups"] = [{"scenes": ["1-1", "1-3"], "reviewed": True, "reason": "林宅客厅到走廊紧接追赶，1-2为另一支线"}]
        book = self.workbook()
        sheet = book["分场表"]
        self.assertEqual(sheet["D2"].border.bottom.style, "thick")
        self.assertEqual(sheet["D3"].border.left.style, "thin")
        self.assertEqual(sheet["D4"].border.top.style, "thick")
        book.close()

    def test_unknown_main_locations_do_not_become_one_shared_group(self) -> None:
        for scene in self.scenes:
            scene["location"] = "房间"
            self.config["scenes"][scene["scene"]].update(main_location="需确认", notes="缺少地点主体，主场景需确认")
        keys = [builder.location_key(scene, self.config) for scene in self.scenes]
        self.assertEqual(len(set(keys)), 6)
        self.assertEqual(builder.validation_issues(self.scenes, self.config, [], strict=True), [])
        self.config["continuity_groups"] = [{"scenes": ["1-1", "1-2"], "reviewed": True, "reason": "待核"}]
        self.assertTrue(any("主场景尚未确认" in x for x in builder.validation_issues(self.scenes, self.config, [], strict=True)))

    def test_invalid_location_metadata_blocks_strict_build(self) -> None:
        for field, value in [("main_location", ""), ("main_location", None), ("location_group", []), ("location_group", "")]:
            with self.subTest(field=field, value=value):
                review = self.config["scenes"]["1-1"]
                old = dict(review)
                review[field] = value
                self.assertTrue(any(field in x for x in builder.validation_issues(self.scenes, self.config, [], strict=True)))
                review.clear()
                review.update(old)


class CastClassificationContractTests(unittest.TestCase):
    def test_reviewed_names_are_not_rejected_by_occupation_or_latin_voice_suffix(self) -> None:
        fixture = VoiceAndVehicleTests()
        fixture.setUp()
        fixture.config["roles"] = [{"name": "李保安", "mark": "李"}, {"name": "Carlos", "mark": "卡"}]
        fixture.review["on_camera"] = ["李保安", "Carlos"]
        roles = builder.resolve_roles([fixture.scene], fixture.config)
        self.assertEqual(builder.validation_issues([fixture.scene], fixture.config, roles, strict=True), [])
        self.assertFalse(builder.is_voice_name("Carlos"))
        self.assertTrue(builder.is_voice_name("Carlos（OS）"))
        self.assertTrue(builder.is_voice_name("Carlos VO"))
        self.assertEqual(builder.collect_voices({"people": ["Carlos"], "body_text": "Carlos：你好。"}), {"vo": "", "os": ""})
        tagged = builder.split_people("人物：Carlos（OS）")
        self.assertEqual(builder.collect_voices({"people": tagged, "body_text": "Carlos（OS）：你好。"}), {"vo": "", "os": "Carlos"})

    def test_named_childhood_and_brief_special_roles_remain_separate_in_both_sheets(self) -> None:
        fixture = VoiceAndVehicleTests()
        fixture.setUp()
        fixture.config["roles"] += [{"name": "童年沈晚晴", "mark": "童"}, {"name": "小周", "mark": "周"}]
        fixture.review.update(on_camera=["童年沈晚晴", "小周"], special_actors="保安甲、护士", groups="其余围观群众", notes="保安甲只有一句台词；护士无台词但单独扶住童年沈晚晴")
        roles = builder.resolve_roles([fixture.scene], fixture.config)
        self.assertEqual(builder.validation_issues([fixture.scene], fixture.config, roles, strict=True), [])
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "cast.xlsx"
            builder.build_workbook(Path("fixture.docx"), [fixture.scene], path, fixture.config)
            book = load_workbook(path)
            for sheet in book:
                cells = {str(sheet.cell(1, c).value).replace("\n", ""): sheet.cell(2, c).value for c in range(1, sheet.max_column + 1)}
                self.assertIsNone(cells["沈晚晴"])
                self.assertEqual(cells["童年沈晚晴"], "童")
                self.assertEqual(cells["小周"], "周")
                self.assertEqual(cells["特约演员"], "保安甲、护士")
                self.assertEqual(cells["群演"], "其余围观群众")
            book.close()

    def test_duplicate_roles_and_multicharacter_marks_are_rejected(self) -> None:
        fixture = VoiceAndVehicleTests()
        fixture.setUp()
        issues = builder.validation_issues([fixture.scene], fixture.config, [("沈晚晴", "晴"), ("沈晚晴", "晚")], strict=True)
        self.assertTrue(any("姓名重复" in x for x in issues))
        issues = builder.validation_issues([fixture.scene], fixture.config, [("沈晚晴", "晚晴")], strict=True)
        self.assertTrue(any("单个字符" in x for x in issues))


if __name__ == "__main__":
    unittest.main()
